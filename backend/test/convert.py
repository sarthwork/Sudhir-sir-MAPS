import pandas as pd
import numpy as np
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Load the dataset
input_file = "convert.xlsx"
output_file = "land_valuation_updated.xlsx"

df = pd.read_excel(input_file)

# 2. Conversion factor
ACRE_TO_SQM = 4046.856

acre_col = "Total Land Area (Acres)"
sqm_col = "Total Land Area (sqm)"

# 3. Coerce to numeric values
df[acre_col] = pd.to_numeric(df[acre_col], errors="coerce")
df[sqm_col] = pd.to_numeric(df[sqm_col], errors="coerce")

# 4. Fill missing values bidirectionally
acre_exists_sqm_missing = df[acre_col].notna() & df[sqm_col].isna()
df.loc[acre_exists_sqm_missing, sqm_col] = (
    df.loc[acre_exists_sqm_missing, acre_col] * ACRE_TO_SQM
)

sqm_exists_acre_missing = df[sqm_col].notna() & df[acre_col].isna()
df.loc[sqm_exists_acre_missing, acre_col] = (
    df.loc[sqm_exists_acre_missing, sqm_col] / ACRE_TO_SQM
)

# 5. Export with formatting
numeric_formats = {
    "Total Land Area (Acres)": "#,##0.00",
    "Total Land Area (sqm)": "#,##0.00",
    "Market Value (INR Mn)": "#,##0.00",
    "M.V. per acre": "#,##0.00",
}

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Valuation")
    ws = writer.sheets["Valuation"]

    # Header row
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 32

    # Body cells + column widths
    for col_idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        is_numeric = col_name in numeric_formats

        # Width = longest value in the column (capped), with a little padding
        longest = max(
            [len(str(col_name))] + [len(str(v)) for v in df[col_name].dropna()]
        )
        ws.column_dimensions[letter].width = min(max(longest + 3, 12), 45)

        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = border
            if is_numeric:
                cell.number_format = numeric_formats[col_name]
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

print(f"Updated successfully! Output saved to: {output_file}")